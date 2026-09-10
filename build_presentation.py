"""
Script to build the professional 14-slide PowerPoint presentation
Seasonal_Agriculture_Performance_Analysis.pptx
Matching the exact structure of VOIS_Major_Project_PPT_Submission_Template.pptx
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE_DIR = r"c:\Users\ragha\.gemini\antigravity-ide\scratch\space-edu\seasonal-agriculture-performance-analysis"
TEMPLATE_PATH = r"C:\Users\ragha\Downloads\Telegram Desktop\VOIS_Major_Project_PPT_Submission_Template.pptx"
OUTPUT_PPTX_PATH = os.path.join(BASE_DIR, "presentation", "Seasonal_Agriculture_Performance_Analysis.pptx")
FIG_DIR = os.path.join(BASE_DIR, "outputs", "figures")

prs = Presentation(TEMPLATE_PATH)
print(f"Loaded template with {len(prs.slides)} slides.")

# Colors
DARK_BLUE = RGBColor(15, 23, 42)
TEXT_GRAY = RGBColor(71, 85, 105)
VOIS_RED = RGBColor(227, 24, 55)
PRIMARY_BLUE = RGBColor(30, 58, 138)

# ------------------------------------------------------------------------------
# SLIDE 1: Title Slide
# ------------------------------------------------------------------------------
slide1 = prs.slides[0]
for shape in slide1.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text
        if "Project Title" in text:
            shape.text_frame.text = "Project Title -\nSeasonal Agriculture Performance Analysis"
            for p in shape.text_frame.paragraphs:
                p.font.bold = True
                p.font.size = Pt(28)
                p.font.color.rgb = DARK_BLUE
        elif "Student Name" in text:
            shape.text_frame.text = "Student Name: Raghav Paroli\nCollege: Dept. of Computer Science & Data Analytics\nVOIS AICTE Batch 2026-2027"
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(15)
                p.font.color.rgb = TEXT_GRAY
        elif "AICTE STU ID" in text:
            shape.text_frame.text = "AICTE STU ID : (its available in your offer letter)"
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(13)
                p.font.bold = True
                p.font.color.rgb = PRIMARY_BLUE

# ------------------------------------------------------------------------------
# SLIDE 2: Problem Statement
# ------------------------------------------------------------------------------
slide2 = prs.slides[1]
# Text is already verbatim from brief; ensure styling is clean
for shape in slide2.shapes:
    if shape.has_text_frame:
        if "PROBLEM" in shape.text_frame.text:
            for p in shape.text_frame.paragraphs:
                p.font.bold = True
                p.font.color.rgb = DARK_BLUE

# ------------------------------------------------------------------------------
# SLIDE 3: Project Description
# ------------------------------------------------------------------------------
slide3 = prs.slides[2]
desc_text = (
    "This project provides an end-to-end data analytics and empirical modeling investigation into 4,000 farm-season records "
    "across 8 Indian states, 8 major crops, and 4 irrigation methods over Kharif, Rabi, and Zaid seasons.\n\n"
    "â€¢ Data Quality & Robust Imputation: Hierarchical group-wise median imputation resolves sensor gaps without collapsing genuine botanical and seasonal climatic variations.\n"
    "â€¢ Comprehensive Feature Engineering: Derived domain metrics (Profit Margin %, Cost/Ha, Revenue/Ha, Fertilizer Intensity, Yield Tertiles) enable standardized multi-crop comparisons.\n"
    "â€¢ Exploratory & Agronomic Profiling: Investigates the 'Monsoon Penalty' paradox, sugarcane yield segmentation, and water efficiency across micro-irrigation regimes.\n"
    "â€¢ Statistical Verification: Employs One-Way ANOVA, Kruskal-Wallis tests, and Tukey HSD post-hoc comparisons to establish high-confidence empirical conclusions.\n"
    "â€¢ Practical Roadmap: Translates statistical findings into actionable recommendations for farmers, agricultural planners, and crop insurance policymakers."
)

# Replace description text
for shape in slide3.shapes:
    if shape.has_text_frame and "Project Description" in shape.text_frame.text:
        shape.text_frame.text = "Project Description"
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(24)
            p.font.color.rgb = DARK_BLUE

# Add structured description text box
tx_box = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(4.8))
tf = tx_box.text_frame
tf.word_wrap = True
tf.text = desc_text
for p in tf.paragraphs:
    p.font.size = Pt(13.5)
    p.font.color.rgb = TEXT_GRAY

# ------------------------------------------------------------------------------
# SLIDE 4: WHO ARE THE END USERS?
# ------------------------------------------------------------------------------
slide4 = prs.slides[3]
end_users_text = (
    "1. Farmers & Cultivators:\n"
    "   Select optimal seasonal crop-mix, transition from flood to micro-irrigation, and preempt Kharif pest surges.\n\n"
    "2. Agricultural Extension Officers & Agronomists:\n"
    "   Deliver data-backed advisories on seed varieties, fertilizer intensities, and integrated pest management (IPM).\n\n"
    "3. Government Policymakers & Planners:\n"
    "   Formulate season-specific Minimum Support Prices (MSP), allocate canal water quotas, and target micro-irrigation subsidies.\n\n"
    "4. Agri-Fintech & Crop Insurance Providers:\n"
    "   Design dynamic actuarial risk premiums aligned with Kharif pest vulnerabilities and seasonal farm cash-flow cycles.\n\n"
    "5. Agrochemical & Farm Input Suppliers:\n"
    "   Optimize seasonal supply chain inventories for NPK fertilizers, micronutrients, and crop protection chemicals.\n\n"
    "6. Agricultural Researchers & Data Scientists:\n"
    "   Serve as an empirical benchmark for climate resilience research, yield forecasting models, and resource sustainability."
)
for shape in slide4.shapes:
    if shape.has_text_frame and "END USERS" in shape.text_frame.text:
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.color.rgb = DARK_BLUE

tx_box = slide4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(4.8))
tf = tx_box.text_frame
tf.word_wrap = True
tf.text = end_users_text
for p in tf.paragraphs:
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_GRAY

# ------------------------------------------------------------------------------
# SLIDE 5: Technology Used
# ------------------------------------------------------------------------------
slide5 = prs.slides[4]
tech_text = (
    "â€¢ Python 3.12: Core programming language for end-to-end data pipeline, modeling, and automation.\n\n"
    "â€¢ Jupyter Notebook: Interactive computational environment for reproducible analysis and narrative storytelling.\n\n"
    "â€¢ Pandas & NumPy: High-performance data manipulation, group-wise imputation, and numerical transformations.\n\n"
    "â€¢ Matplotlib & Seaborn: Publication-grade 300 DPI visualizations, seasonal multi-panel plots, and annotated heatmaps.\n\n"
    "â€¢ SciPy & Statsmodels: Rigorous inferential statisticsâ€”One-Way ANOVA, Kruskal-Wallis, Tukey HSD, and Chi-Square tests.\n\n"
    "â€¢ Plotly: Dynamic interactive exploratory visualization with multi-dimensional filtering.\n\n"
    "â€¢ Git & GitHub: Professional version control, collaborative code management, and open-source project dissemination."
)
for shape in slide5.shapes:
    if shape.has_text_frame and "Technology" in shape.text_frame.text:
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.color.rgb = DARK_BLUE

tx_box = slide5.shapes.add_textbox(Inches(2.5), Inches(1.8), Inches(9.8), Inches(4.8))
tf = tx_box.text_frame
tf.word_wrap = True
tf.text = tech_text
for p in tf.paragraphs:
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_GRAY

# ------------------------------------------------------------------------------
# SLIDE 6: RESULTS - 1: Seasonal Yield & Profitability Dynamics
# ------------------------------------------------------------------------------
slide6 = prs.slides[5]
for shape in slide6.shapes:
    if shape.has_text_frame and "RESULTS" in shape.text_frame.text:
        shape.text_frame.text = "RESULTS: Seasonal Yield & Farm Profitability Dynamics"
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(22)
            p.font.color.rgb = DARK_BLUE

# Add chart
img_path = os.path.join(FIG_DIR, "fig04_seasonal_profitability_ci.png")
if os.path.exists(img_path):
    slide6.shapes.add_picture(img_path, Inches(0.8), Inches(1.7), width=Inches(7.2))

# Add commentary
tx_box = slide6.shapes.add_textbox(Inches(8.2), Inches(1.8), Inches(4.5), Inches(4.5))
tf = tx_box.text_frame
tf.word_wrap = True
tf.text = (
    "Key Takeaways & Findings:\n\n"
    "â€¢ Zaid Outperformance: Achieves the highest average net profit (INR 232,042) and profit margin (11.5%), with 58.6% profitable farms.\n\n"
    "â€¢ Kharif Margin Squeeze: Generates only INR 28,323 mean profit, with 53.5% of farms operating at a financial loss due to elevated input costs.\n\n"
    "â€¢ Controlled Rabi Stability: Delivers steady profits (INR 119,548; 52.2% profitable farms), proving managed irrigation mitigates climate risks."
)
for p in tf.paragraphs:
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_GRAY

# ------------------------------------------------------------------------------
# SLIDE 7: RESULTS - 2: Crop x Season Profitability Matrix
# ------------------------------------------------------------------------------
slide7 = prs.slides[6]
for shape in slide7.shapes:
    if shape.has_text_frame and "RESULTS" in shape.text_frame.text:
        shape.text_frame.text = "RESULTS: Crop Ã— Season Profitability Matrix"
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(22)
            p.font.color.rgb = DARK_BLUE
    elif shape.has_text_frame and "[Add screen shots" in shape.text_frame.text:
        shape.text_frame.text = ""

img_path = os.path.join(FIG_DIR, "fig13_crop_season_profitability_heatmap.png")
if os.path.exists(img_path):
    slide7.shapes.add_picture(img_path, Inches(0.8), Inches(1.7), width=Inches(6.8))

tx_box = slide7.shapes.add_textbox(Inches(7.8), Inches(1.8), Inches(4.8), Inches(4.5))
tf = tx_box.text_frame
tf.word_wrap = True
tf.text = (
    "Key Takeaways & Findings:\n\n"
    "â€¢ Commercial Crop Superiority: Chilli, Sugarcane, and Groundnut yield positive profits across all seasons, peaking in Rabi and Zaid.\n\n"
    "â€¢ Grain Crop Vulnerability: Rice and Maize experience negative net margins during Kharif, where high seed and pesticide costs outpace farm-gate prices.\n\n"
    "â€¢ Seasonal Specialization: Wheat achieves optimal profitability in Rabi (+INR 134k) versus Kharif (-INR 46k), highlighting the critical role of temperature timing."
)
for p in tf.paragraphs:
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_GRAY

# ------------------------------------------------------------------------------
# SLIDE 8: RESULTS - 3: Environmental Drivers & Disease/Pest Vulnerability
# ------------------------------------------------------------------------------
slide8 = prs.slides[7]
for shape in slide8.shapes:
    if shape.has_text_frame and "RESULTS" in shape.text_frame.text:
        shape.text_frame.text = "RESULTS: Environmental Drivers & Pest/Disease Vulnerability"
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(22)
            p.font.color.rgb = DARK_BLUE
    elif shape.has_text_frame and "[Add screen shots" in shape.text_frame.text:
        shape.text_frame.text = ""

img_path = os.path.join(FIG_DIR, "fig07_disease_pest_risk_by_season.png")
if os.path.exists(img_path):
    slide8.shapes.add_picture(img_path, Inches(0.8), Inches(1.7), width=Inches(7.2))

tx_box = slide8.shapes.add_textbox(Inches(8.2), Inches(1.8), Inches(4.5), Inches(4.5))
tf = tx_box.text_frame
tf.word_wrap = True
tf.text = (
    "Key Takeaways & Findings:\n\n"
    "â€¢ Monsoon Pest Surge: Kharif disease and pest risk averages 52.88%, compared to 38.89% in Rabi and 36.63% in Zaid.\n\n"
    "â€¢ High Effect Size: One-Way ANOVA proves this disparity is extraordinary (F = 1049.47, p < 10^-300, Î·Â² = 0.3443).\n\n"
    "â€¢ Moisture-Infestation Link: Strong linear correlation between atmospheric humidity (>70%) and pest risk drives a 25% surge in pesticide expenditure."
)
for p in tf.paragraphs:
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_GRAY

# ------------------------------------------------------------------------------
# SLIDE 9: RESULTS - 4: Irrigation Method Effectiveness Across Seasons
# ------------------------------------------------------------------------------
slide9 = prs.slides[8]
for shape in slide9.shapes:
    if shape.has_text_frame and "RESULTS" in shape.text_frame.text:
        shape.text_frame.text = "RESULTS: Irrigation Method Effectiveness Across Seasons"
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(22)
            p.font.color.rgb = DARK_BLUE
    elif shape.has_text_frame and "[Add screen shots" in shape.text_frame.text:
        shape.text_frame.text = ""

img_path = os.path.join(FIG_DIR, "fig06_water_efficiency_irrigation_heatmap.png")
if os.path.exists(img_path):
    slide9.shapes.add_picture(img_path, Inches(0.8), Inches(1.7), width=Inches(7.2))

tx_box = slide9.shapes.add_textbox(Inches(8.2), Inches(1.8), Inches(4.5), Inches(4.5))
tf = tx_box.text_frame
tf.word_wrap = True
tf.text = (
    "Key Takeaways & Findings:\n\n"
    "â€¢ Micro-Irrigation Superiority: Drip irrigation achieves 4.82 t/1000mÂ³ water efficiency, beating Flood systems (1.64 t/1000mÂ³) by nearly 3x.\n\n"
    "â€¢ Flood Inefficiency: Flood irrigation consumes >9,000 mÂ³ of water on average, elevating pumping electricity and extraction costs.\n\n"
    "â€¢ Summer Water Optimization: Zaid micro-irrigation maximizes yield per unit water applied during peak evaporation conditions."
)
for p in tf.paragraphs:
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_GRAY

# ------------------------------------------------------------------------------
# SLIDE 10: RESULTS - 5: Statistical Hypothesis Testing Summary
# ------------------------------------------------------------------------------
slide10 = prs.slides[9]
for shape in slide10.shapes:
    if shape.has_text_frame and "RESULTS" in shape.text_frame.text:
        shape.text_frame.text = "RESULTS: Statistical Hypothesis Testing Summary"
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(22)
            p.font.color.rgb = DARK_BLUE
    elif shape.has_text_frame and "[Add screen shots" in shape.text_frame.text:
        shape.text_frame.text = ""

img_path = os.path.join(FIG_DIR, "fig12_statistical_tests_summary_table.png")
if os.path.exists(img_path):
    slide10.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), width=Inches(11.5))

# ------------------------------------------------------------------------------
# SLIDE 11: Future scope
# ------------------------------------------------------------------------------
slide11 = prs.slides[10]
future_text = (
    "â€¢ Multi-Year Longitudinal Modeling:\n"
    "  Expand the cross-sectional dataset across multi-decadal time-series to capture climate change impacts and El NiÃ±o cycles.\n\n"
    "â€¢ Predictive Yield & Revenue Machine Learning:\n"
    "  Deploy ensemble models (XGBoost, LightGBM, CatBoost) to forecast pre-sowing farm profits based on early-season meteorology.\n\n"
    "â€¢ Remote Sensing & Satellite NDVI Integration:\n"
    "  Incorporate Sentinel-2 and Landsat multispectral imagery to validate soil moisture, crop canopy vigor, and field stress remotely.\n\n"
    "â€¢ Hyper-Local IoT Agro-Advisory System:\n"
    "  Deploy low-cost farm IoT sensor nodes for micro-climatic pest alerts, dynamic irrigation scheduling, and mobile advisories."
)
for shape in slide11.shapes:
    if shape.has_text_frame and "Future scope" in shape.text_frame.text:
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.color.rgb = DARK_BLUE

tx_box = slide11.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(4.8))
tf = tx_box.text_frame
tf.word_wrap = True
tf.text = future_text
for p in tf.paragraphs:
    p.font.size = Pt(13.5)
    p.font.color.rgb = TEXT_GRAY

# ------------------------------------------------------------------------------
# SLIDE 12: GitHub Link
# ------------------------------------------------------------------------------
slide12 = prs.slides[11]
repo_info = (
    "\n\nRepository Contents & Deliverables:\n"
    "â€¢ Full Executed Jupyter Notebook (47 cells, complete outputs, visualizations, markdown)\n"
    "â€¢ Raw and Cleaned Processed Datasets (data/raw & data/processed)\n"
    "â€¢ 15 High-Resolution 300 DPI Publication-Grade Figures (outputs/figures)\n"
    "â€¢ Complete 14-Slide Official PPTX Presentation Deck (presentation/)\n"
    "â€¢ Comprehensive Technical README and Environment Requirements"
)
for shape in slide12.shapes:
    if shape.has_text_frame:
        if "GitHub Link" in shape.text_frame.text:
            for p in shape.text_frame.paragraphs:
                p.font.bold = True
                p.font.color.rgb = DARK_BLUE
        elif "https://github.com" in shape.text_frame.text:
            shape.text_frame.text = "https://github.com/raghavparoli07-art/seasonal-agriculture-performance-analysis" + repo_info
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(13)
                p.font.color.rgb = TEXT_GRAY

# ------------------------------------------------------------------------------
# SLIDE 13: VOIS Course completion certificate
# ------------------------------------------------------------------------------
slide13 = prs.slides[12]
for shape in slide13.shapes:
    if shape.has_text_frame and "completion certificate" in shape.text_frame.text:
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.color.rgb = DARK_BLUE

cert_img = os.path.join(FIG_DIR, "vois_data_viz_certificate.png")
if os.path.exists(cert_img):
    slide13.shapes.add_picture(cert_img, Inches(1.8), Inches(1.6), width=Inches(8.8))

# ------------------------------------------------------------------------------
# SLIDE 14: Thank you
# ------------------------------------------------------------------------------
slide14 = prs.slides[13]
for shape in slide14.shapes:
    if shape.has_text_frame and "Thank you" in shape.text_frame.text:
        shape.text_frame.text = "Thank You\nQuestions & Discussion"
        for p in shape.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(36)
            p.font.color.rgb = DARK_BLUE

tx_box = slide14.shapes.add_textbox(Inches(1.5), Inches(3.5), Inches(9.5), Inches(2.0))
tf = tx_box.text_frame
tf.word_wrap = True
tf.text = (
    "Seasonal Agriculture Performance Analysis\n"
    "VOIS AICTE Batch 2026-2027 Major Project\n"
    "Author: Raghav Paroli  •  GitHub: github.com/raghavparoli07-art/seasonal-agriculture-performance-analysis"
)
for p in tf.paragraphs:
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_GRAY

# Save presentation
prs.save(OUTPUT_PPTX_PATH)
print(f"Presentation successfully created and saved to: {OUTPUT_PPTX_PATH}")

